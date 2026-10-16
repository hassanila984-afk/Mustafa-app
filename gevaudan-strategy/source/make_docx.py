# -*- coding: utf-8 -*-
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from content import DOC

NAVY = RGBColor(0x0E, 0x17, 0x26)
GOLD = RGBColor(0xA8, 0x84, 0x3E)
GREY = RGBColor(0x55, 0x5B, 0x66)
FONT = "Arial"

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(2)
    # اتجاه القسم من اليمين لليسار
    s._sectPr.append(OxmlElement("w:bidi"))

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(12)
rpr = st.element.get_or_add_rPr()
rf = rpr.find(qn("w:rFonts"))
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rf.set(qn(a), FONT)
szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), "24"); rpr.append(szcs)


def rtl_par(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    ppr = p._p.get_or_add_pPr()
    ppr.insert(0, OxmlElement("w:bidi"))
    p.alignment = align
    return p


def run(p, text, size=12, bold=False, color=None, italic=False):
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.rtl = True
    rp = r._r.get_or_add_rPr()
    if bold:
        rp.append(OxmlElement("w:bCs"))
    sc = OxmlElement("w:szCs"); sc.set(qn("w:val"), str(int(size * 2))); rp.append(sc)
    if color is not None:
        r.font.color.rgb = color
    return r


def spacing(p, before=0, after=6, line=1.35):
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = line


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def borders(table, color="D9CBA8"):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), color)
        b.append(el)
    tblPr.append(b)


def add_table(rows):
    ncol = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t._tbl.tblPr.append(OxmlElement("w:bidiVisual"))
    borders(t)
    first_col_label = rows[0][0] == ""
    for i, row in enumerate(rows):
        trPr = t.rows[i]._tr.get_or_add_trPr(); trPr.append(OxmlElement("w:cantSplit"))
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            p = rtl_par(c.paragraphs[0]); spacing(p, 2, 2, 1.2)
            header = i == 0 and any(rows[0])
            label = j == 0 and first_col_label and i > 0
            if header:
                shade(c, "0E1726"); run(p, val, 11, True, RGBColor(0xE8, 0xD5, 0xA9))
            elif label:
                shade(c, "F3EEE2"); run(p, val, 11, True, NAVY)
            else:
                if i % 2 == 0:
                    shade(c, "FAF8F3")
                run(p, val, 11)
    if not any(rows[0]):  # صف عنوان فارغ: احذفه
        t._tbl.remove(t.rows[0]._tr)
    doc.add_paragraph()


for kind, data in DOC:
    if kind == "title":
        p = rtl_par(doc.add_paragraph()); spacing(p, 60, 6)
        run(p, data, 30, True, NAVY)
    elif kind == "subtitle":
        p = rtl_par(doc.add_paragraph()); spacing(p, 0, 6)
        run(p, data, 15, False, GOLD)
    elif kind == "meta":
        p = rtl_par(doc.add_paragraph()); spacing(p, 0, 30)
        run(p, data, 11, False, GREY)
    elif kind == "h1":
        p = rtl_par(doc.add_paragraph()); spacing(p, 18, 8)
        p.paragraph_format.keep_with_next = True
        run(p, data, 18, True, NAVY)
    elif kind == "h2":
        p = rtl_par(doc.add_paragraph()); spacing(p, 10, 4)
        p.paragraph_format.keep_with_next = True
        run(p, data, 13.5, True, GOLD)
    elif kind == "p":
        p = rtl_par(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.JUSTIFY); spacing(p)
        run(p, data)
    elif kind == "quote":
        p = rtl_par(doc.add_paragraph()); spacing(p, 6, 10)
        p.paragraph_format.right_indent = Cm(0.8); p.paragraph_format.left_indent = Cm(0.8)
        ppr = p._p.get_or_add_pPr()
        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "F3EEE2"); ppr.append(sh)
        run(p, data, 13, True, NAVY)
    elif kind in ("bullets", "numbers"):
        for i, item in enumerate(data, 1):
            p = rtl_par(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.JUSTIFY); spacing(p, 0, 4)
            p.paragraph_format.right_indent = Cm(0.6)
            mark = f"{i}.  " if kind == "numbers" else "◆  "
            run(p, mark, 11, True, GOLD)
            if "：" in item or ": " in item[:40]:
                head, rest = item.split(": ", 1)
                run(p, head + ": ", 12, True, NAVY); run(p, rest)
            else:
                run(p, item)
    elif kind == "table":
        add_table(data)

# ترقيم الصفحات في التذييل
for s in doc.sections:
    fp = rtl_par(s.footer.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
    run(fp, "جيفودان العراق | الاستراتيجية  ·  ", 9, False, GREY)
    r = fp.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); r._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.text = txt; r._r.append(it)
    r.font.size = Pt(9)

doc.save(sys.argv[1])
print("saved", sys.argv[1])
